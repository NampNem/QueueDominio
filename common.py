# ============================================================
# NFS-e (relatório do portal) -> Excel + TXT para importação no Domínio
# Arquivo: nfse_dominio.py  (mesma pasta do app principal)
# ============================================================
import io
import itertools
import os

import pandas as pd
import streamlit as st

ARQ_BANCO = "banco_tributacao_dominio.xlsx"   # banco: Cód. Tributação Nacional -> Conta Débito

# Colunas do relatório (nome exato no cabeçalho)
COL = {
    "num": "Número NFS-e",
    "data": "Data Geração",
    "cnpj": "CNPJ/CPF Prestador",
    "forn": "Nome Prestador",
    "bruto": "Valor do Serviço (R$)",
    "iss": "Valor do ISSQN (R$)",
    "ret_iss": "Retenção ISSQN",
    "trib": "Cód. Tributação Nacional",      # coluna U
    "desc": "Descrição do Serviço",
    "irrf": "IRRF (R$)",
    "pcc": "Contrib. Sociais Ret. (R$)",       # PIS+COFINS+CSLL retidos (unificado)
    "inss": "Contrib. Previd. Ret. (R$)",
    "sit": "Situação",
}

# Ordem dos lançamentos de retenção no múltiplo
ORDEM_RET = ["IRRF", "INSS", "ISS", "PCC"]
NOME_RET = {"IRRF": "IRRF", "INSS": "INSS", "ISS": "ISS", "PCC": "PIS/COFINS/CSLL"}


# ------------------------------------------------------------
# Utilidades
# ------------------------------------------------------------
def _num(x):
    try:
        v = float(x)
        return 0.0 if pd.isna(v) else round(v, 2)
    except (TypeError, ValueError):
        return 0.0


def _br(v):
    """4800.5 -> '4800,50' (sem separador de milhar)"""
    return f"{v:.2f}".replace(".", ",")


def _limpa(txt):
    return str(txt).replace(";", " ").replace("\n", " ").replace("\r", " ").strip()


def _fmt_doc(x):
    try:
        d = str(int(float(x)))
    except (TypeError, ValueError):
        return ""
    return d.zfill(14) if len(d) > 11 else d.zfill(11)


# ------------------------------------------------------------
# Banco de dados (Cód. Tributação -> Conta Débito)
# ------------------------------------------------------------
def carregar_banco():
    if os.path.exists(ARQ_BANCO):
        try:
            df = pd.read_excel(ARQ_BANCO, dtype=str).fillna("")
            for c in ["Código", "Descrição", "Conta Débito"]:
                if c not in df.columns:
                    df[c] = ""
            return df[["Código", "Descrição", "Conta Débito"]]
        except Exception:
            pass
    return pd.DataFrame(columns=["Código", "Descrição", "Conta Débito"])


def salvar_banco(df):
    df.to_excel(ARQ_BANCO, index=False)


def mesclar_banco(banco, novos):
    """Acrescenta ao banco os códigos que ainda não existem."""
    existentes = set(banco["Código"].astype(str))
    add = novos[~novos["Código"].isin(existentes)]
    if add.empty:
        return banco
    add = add.assign(**{"Conta Débito": ""})[["Código", "Descrição", "Conta Débito"]]
    return pd.concat([banco, add], ignore_index=True)


# ------------------------------------------------------------
# Retenções: testa todas as combinações (líquido + retenções = bruto)
# ------------------------------------------------------------
def achar_combinacoes(bruto, liquido, candidatos):
    """candidatos: {'IRRF': 72.0, 'INSS': 0.., ...} (só valores > 0).
    Devolve as combinações cuja soma = bruto - líquido."""
    dif = round(bruto - liquido, 2)
    if abs(dif) < 0.005:
        return [[]]
    nomes = list(candidatos)
    achados = []
    for r in range(1, len(nomes) + 1):
        for comb in itertools.combinations(nomes, r):
            if abs(sum(candidatos[n] for n in comb) - dif) <= 0.01:
                achados.append(list(comb))
    return achados


def processar_relatorio(arquivo):
    xl = pd.ExcelFile(arquivo)
    aba = "Relação" if "Relação" in xl.sheet_names else xl.sheet_names[0]
    df = xl.parse(aba)
    df.columns = [str(c).strip() for c in df.columns]

    if COL["trib"] not in df.columns and len(df.columns) > 20:
        COL["trib"] = df.columns[20]  # coluna U
    faltando = [v for k, v in COL.items() if v not in df.columns]
    if faltando:
        raise ValueError("Colunas não encontradas no relatório: " + ", ".join(faltando))

    # coluna de valor líquido (se o relatório tiver uma)
    col_liq = next((c for c in df.columns if "líquid" in c.lower() or "liquid" in c.lower()), None)

    df = df[df[COL["num"]].notna()].copy()           # remove linha de TOTAL
    total = len(df)
    df = df[df[COL["sit"]].astype(str).str.strip().str.lower() == "normal"].copy()
    ignoradas = total - len(df)                      # canceladas / substituídas

    registros = []
    for _, r in df.iterrows():
        bruto = _num(r[COL["bruto"]])
        cand = {
            "IRRF": _num(r[COL["irrf"]]),
            "INSS": _num(r[COL["inss"]]),
            "ISS": _num(r[COL["iss"]]) if str(r[COL["ret_iss"]]).strip().startswith("2") else 0.0,
            "PCC": _num(r[COL["pcc"]]),
        }
        cand = {k: v for k, v in cand.items() if v > 0}
        obs = ""

        if col_liq is not None and pd.notna(r[col_liq]):
            liquido = _num(r[col_liq])
            combos = achar_combinacoes(bruto, liquido, cand)
            if combos:
                # prefere a combinação igual ao que o relatório informa
                esc = next((c for c in combos if set(c) == set(cand)), combos[0])
                ret = {k: cand[k] for k in esc}
                obs = "Retenções localizadas pelo líquido"
            else:
                ret = cand
                obs = "ATENÇÃO: líquido não fecha com nenhuma combinação"
        else:
            ret = cand
            liquido = round(bruto - sum(ret.values()), 2)

        cod, _, nome_trib = str(r[COL["trib"]]).partition(" - ")
        registros.append({
            "Data": pd.to_datetime(r[COL["data"]]),
            "Nº NF": int(float(r[COL["num"]])),
            "Fornecedor": _limpa(r[COL["forn"]]),
            "CNPJ/CPF": _fmt_doc(r[COL["cnpj"]]),
            "Cód. Tributação": cod.strip(),
            "Tributação (descrição)": nome_trib.strip(),
            "Descrição do Serviço": _limpa(r[COL["desc"]]),
            "Valor Bruto": bruto,
            "IRRF": ret.get("IRRF", 0.0),
            "INSS": ret.get("INSS", 0.0),
            "ISS Retido": ret.get("ISS", 0.0),
            "PIS/COFINS/CSLL": ret.get("PCC", 0.0),
            "Valor Líquido": liquido,
            "Retenções": ", ".join(NOME_RET[k] for k in ORDEM_RET if k in ret) or "Sem retenção",
            "Obs": obs,
        })

    res = pd.DataFrame(registros).sort_values(["Data", "Nº NF"]).reset_index(drop=True)
    return res, ignoradas, col_liq is not None


# ------------------------------------------------------------
# Lançamentos Domínio
# ------------------------------------------------------------
def montar_lancamentos(notas, banco, contas, matriz, cc_deb, cc_cred, historico):
    """Devolve (linhas, notas_sem_conta)."""
    mapa = {str(c): str(v).strip() for c, v in zip(banco["Código"], banco["Conta Débito"])}
    linhas, sem_conta = [], []
    lote = 0

    def linha(data, deb, cred, valor, compl, lote_n):
        return [data, deb, cred, _br(valor), historico, compl, lote_n, matriz, cc_deb, cc_cred]

    for n in notas.itertuples(index=False):
        data = n.Data.strftime("%d/%m/%Y")
        forn = n.Fornecedor
        conta_deb = mapa.get(n._4, "")
        if not conta_deb:
            sem_conta.append(f"NF {n._1} - {forn} (cód. {n._4})")

        rets = [
            ("IRRF", n.IRRF), ("INSS", n.INSS),
            ("ISS", n._10), ("PCC", n._11),
        ]
        rets = [(k, v) for k, v in rets if v > 0]
        lote += 1
        compl_nf = f"NF - {n._1} - {forn}"

        if not rets:  # lançamento simples
            linhas.append(linha(data, conta_deb, contas["fornecedor"], n._7, compl_nf, lote))
            continue

        # múltiplo: bruto só no débito; demais linhas sem lote
        linhas.append(linha(data, conta_deb, "", n._7, compl_nf, lote))
        for k, v in rets:
            compl = f"Retenção {NOME_RET[k]} s/ NF {n._1} - {forn}"
            linhas.append(linha(data, "", contas[k], v, compl, ""))
        linhas.append(linha(data, "", contas["fornecedor"], n._12, compl_nf, ""))

    return linhas, sem_conta


def gerar_txt(linhas, cabecalho, encoding):
    cab = ["Data", "Cód. Conta Debito", "Cód. Conta Credito", "Valor", "Cód. Histórico",
           "Complemento Histórico", "Inicia Lote", "Código Matriz/Filial",
           "Centro de Custo Débito", "Centro de Custo Crédito"]
    out = []
    if cabecalho:
        out.append(";".join(cab))
    for l in linhas:
        out.append(";".join(_limpa(x) for x in l))
    return ("\r\n".join(out) + "\r\n").encode(encoding, errors="replace")


def gerar_excel(notas):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
        notas.to_excel(w, index=False, sheet_name="NFS-e")
        ws = w.sheets["NFS-e"]
        from openpyxl.styles import Font, PatternFill
        for c in ws[1]:
            c.font = Font(bold=True, name="Arial", color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        larg = {"A": 12, "B": 10, "C": 38, "D": 18, "E": 12, "F": 40, "G": 50,
                "H": 14, "I": 12, "J": 12, "K": 12, "L": 16, "M": 14, "N": 28, "O": 40}
        for col, wd in larg.items():
            ws.column_dimensions[col].width = wd
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial")
                if c.column in range(8, 14):
                    c.number_format = "#,##0.00"
        ws.freeze_panes = "A2"
    return buf.getvalue()


# ------------------------------------------------------------
# Página Streamlit
# ------------------------------------------------------------
def pagina_nfse_dominio():
    if st.button("⬅️ Voltar ao menu"):
        st.session_state["pagina"] = "menu"
        st.rerun()

    st.title("🧾 NFS-e (relatório) → Excel + TXT Domínio")

    # ---- configurações
    with st.expander("⚙️ Contas e parâmetros", expanded=False):
        c1, c2, c3 = st.columns(3)
        contas = {
            "fornecedor": c1.text_input("Conta crédito fornecedor (líquido)", "325"),
            "IRRF": c2.text_input("Conta IRRF", "178"),
            "INSS": c3.text_input("Conta INSS", "184"),
            "ISS": c1.text_input("Conta ISS", "183"),
            "PCC": c2.text_input("Conta PIS/COFINS/CSLL", "3924"),
        }
        historico = c3.text_input("Cód. Histórico", "")
        matriz = c1.text_input("Código Matriz/Filial", "1")
        cc_deb = c2.text_input("Centro de custo débito", "")
        cc_cred = c3.text_input("Centro de custo crédito", "")
        cabecalho = c1.checkbox("Incluir cabeçalho no TXT", value=False)
        encoding = c2.selectbox("Codificação do TXT", ["cp1252", "utf-8", "latin-1"])

    # ---- banco
    if "banco_nfse" not in st.session_state:
        st.session_state["banco_nfse"] = carregar_banco()

    st.subheader("1. Relatório de NFS-e")
    arquivo = st.file_uploader("Envie o relatório (.xlsx)", type=["xlsx"])

    notas = None
    if arquivo is not None:
        try:
            notas, ignoradas, achou_liq = processar_relatorio(arquivo)
        except Exception as e:
            st.error(f"Não consegui ler o relatório: {e}")
            return
        msg = f"{len(notas)} notas lidas"
        if ignoradas:
            msg += f" ({ignoradas} canceladas/substituídas ignoradas)"
        st.success(msg)
        if not achou_liq:
            st.info("O relatório não tem coluna de valor líquido: as retenções foram lidas "
                    "das colunas IRRF, Contrib. Previd., Contrib. Sociais e ISS retido, e o líquido = bruto − retenções.")

        novos = (notas[["Cód. Tributação", "Tributação (descrição)"]]
                 .drop_duplicates("Cód. Tributação")
                 .rename(columns={"Cód. Tributação": "Código", "Tributação (descrição)": "Descrição"}))
        st.session_state["banco_nfse"] = mesclar_banco(st.session_state["banco_nfse"], novos)

    st.subheader("2. Banco: Cód. Tributação Nacional → Conta Débito")
    edit = st.data_editor(
        st.session_state["banco_nfse"],
        num_rows="dynamic",
        use_container_width=True,
        key="editor_banco_nfse",
    )
    st.session_state["banco_nfse"] = edit.fillna("").astype(str)

    b1, b2 = st.columns(2)
    if b1.button("💾 Salvar banco"):
        try:
            salvar_banco(st.session_state["banco_nfse"])
            st.success(f"Banco salvo em {ARQ_BANCO}")
        except Exception as e:
            st.error(f"Erro ao salvar: {e}")
    buf = io.BytesIO()
    st.session_state["banco_nfse"].to_excel(buf, index=False)
    b2.download_button("⬇️ Baixar banco (.xlsx)", buf.getvalue(), file_name=ARQ_BANCO)

    if notas is None:
        return

    # ---- gerar
    st.subheader("3. Gerar arquivos")
    if st.button("🚀 Gerar Excel + TXT", type="primary"):
        linhas, sem_conta = montar_lancamentos(
            notas, st.session_state["banco_nfse"], contas, matriz, cc_deb, cc_cred, historico)
        st.session_state["res_nfse"] = {
            "excel": gerar_excel(notas),
            "txt": gerar_txt(linhas, cabecalho, encoding),
            "qtd": len(linhas),
            "sem_conta": sem_conta,
        }

    res = st.session_state.get("res_nfse")
    if res:
        if res["sem_conta"]:
            st.warning("Notas sem conta débito no banco (ficaram em branco no TXT):\n\n- "
                       + "\n- ".join(res["sem_conta"]))
        st.write(f"{res['qtd']} linhas de lançamento geradas.")
        d1, d2 = st.columns(2)
        d1.download_button("📊 Baixar Excel", res["excel"], file_name="nfse_resumo.xlsx")
        d2.download_button("📄 Baixar TXT Domínio", res["txt"], file_name="lancamentos_dominio.txt")
        st.dataframe(notas, use_container_width=True)
